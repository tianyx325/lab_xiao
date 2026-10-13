# -*- coding: utf-8 -*-
"""Import projects and patents from Excel into Hugo content Markdown.

Sources (default):
  D:\\Claude code\\theme-research-group\\项目清单1.xlsx
  D:\\Claude code\\theme-research-group\\专利清单.xlsx

Source of truth for bulk sync: Excel. Day-to-day edits: Markdown (re-import overwrites).
Upserts by project_code / patent_number. Skips funding_amount in page body.
Only imports rows with is_active truthy (missing treated as active).

Usage:
  py -3 scripts/import_projects_patents.py
  py -3 scripts/import_projects_patents.py --dry-run
"""
from __future__ import annotations

import argparse
import re
import shutil
from datetime import datetime, date
from pathlib import Path

import openpyxl
from pypinyin import lazy_pinyin

ROOT = Path(__file__).resolve().parents[1]
# Defaults; overridden in main() via CLI
XLSX_DIR = ROOT.parent
PROJECT_XLSX = XLSX_DIR / "项目清单1.xlsx"
PATENT_XLSX = XLSX_DIR / "专利清单.xlsx"
CONTENT_DIR = ROOT / "content"
PROJECT_DIR = CONTENT_DIR / "project"
PATENT_DIR = CONTENT_DIR / "patent"


def truthy(v) -> bool:
    if v is None or v == "":
        return True
    if isinstance(v, bool):
        return v
    return str(v).strip().lower() in {"1", "true", "yes", "y", "是"}


def as_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def as_date(v) -> str:
    s = as_str(v)
    if not s:
        return ""
    # "2024-01-01 00:00:00" → date only
    if " " in s:
        s = s.split(" ", 1)[0]
    return s


def ym(s: str) -> str:
    """2024-01-01 → 2024/01"""
    if not s or len(s) < 7:
        return s or "—"
    return s[:7].replace("-", "/")


def slugify_title(title: str, prefix: str = "item") -> str:
    parts = lazy_pinyin(title or "")
    raw = "-".join(parts).lower()
    raw = re.sub(r"[^a-z0-9]+", "-", raw).strip("-")
    if not raw:
        raw = prefix
    return raw[:60].strip("-")


def yaml_quote(s: str) -> str:
    s = (s or "").replace("\\", "\\\\").replace('"', '\\"')
    return f'"{s}"'


def yaml_list(items: list[str], indent: int = 2) -> str:
    sp = " " * indent
    return "\n".join(f"{sp}- {it}" for it in items)


def load_rows(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None or str(row[0]).strip() == "":
            continue
        d = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            d[str(h).strip()] = row[i] if i < len(row) else None
        if not truthy(d.get("is_active")):
            continue
        rows.append(d)
    return rows


def status_zh_project(status: str) -> str:
    s = (status or "").lower()
    return {"ongoing": "在研", "completed": "已结题", "pending": "筹备中"}.get(s, status or "—")


def status_zh_patent(status: str) -> str:
    s = (status or "").lower()
    return {
        "granted": "已授权",
        "pending": "申请中",
        "published": "已公开",
        "rejected": "驳回",
    }.get(s, status or "—")


def funding_tags(source: str, area: str) -> list[str]:
    tags: list[str] = []
    src = as_str(source)
    if "国家自然科学基金" in src:
        tags.append("国家自然科学基金")
        if "重点" in src:
            tags.append("重点项目")
        elif "面上" in src:
            tags.append("面上项目")
    elif "省" in src and "自然" in src:
        tags.append("省级基金")
    elif src:
        # short label from source
        short = re.split(r"[（(]", src, 1)[0].strip()
        if short:
            tags.append(short[:20])

    for part in re.split(r"[；;、，,]", as_str(area)):
        part = part.strip()
        if part and part not in tags:
            tags.append(part)
    # de-dupe preserve order, max ~6
    seen = set()
    out = []
    for t in tags:
        if t not in seen:
            seen.add(t)
            out.append(t)
        if len(out) >= 6:
            break
    return out or ["科研项目"]


def project_slug(row: dict) -> str:
    code = as_str(row.get("project_code"))
    # openpyxl may give int
    if code.endswith(".0"):
        code = code[:-2]
    if code:
        # NSFC-looking numeric codes
        if code.isdigit() or re.match(r"^\d+$", code):
            return f"proj-{code}"
        safe = re.sub(r"[^a-zA-Z0-9\-]+", "-", code).strip("-").lower()
        return f"proj-{safe}" if safe else slugify_title(as_str(row.get("title")), "proj")
    return slugify_title(as_str(row.get("title")), "proj")


def patent_slug(row: dict) -> str:
    num = as_str(row.get("patent_number")).upper()
    digits = re.sub(r"[^0-9]", "", num)
    if digits:
        return f"zl-{digits}"
    return slugify_title(as_str(row.get("title")), "patent")


def render_project(row: dict) -> str:
    """Return full markdown file contents for a project row."""
    # reuse write_project logic via temp — build string directly
    title = as_str(row.get("title"))
    summary = as_str(row.get("short_description")) or as_str(row.get("description"))[:80]
    start = as_date(row.get("start_date"))
    end = as_date(row.get("actual_end_date")) or as_date(row.get("expected_end_date"))
    code = as_str(row.get("project_code"))
    if code.endswith(".0"):
        code = code[:-2]
    source = as_str(row.get("funding_source"))
    pi = as_str(row.get("principal_investigator")) or "肖照林"
    status = status_zh_project(as_str(row.get("status")))
    desc = as_str(row.get("description"))
    tags = funding_tags(source, as_str(row.get("research_area")))
    authors = ["admin"] if "肖照林" in pi else []
    period = f"{ym(start)} – {ym(end)}" if start or end else "—"

    body_lines = [f"**项目类型**：{source or '—'}  "]
    if code:
        body_lines.append(f"**项目编号**：{code}  ")
    body_lines.extend(
        [
            f"**起止时间**：{period}  ",
            f"**负责人**：{pi}  ",
            f"**状态**：{status}  ",
            "",
        ]
    )
    if desc:
        body_lines.extend(["### 研究内容", "", desc])

    fm = [
        "---",
        f"title: {yaml_quote(title)}",
        f"summary: {yaml_quote(summary)}",
        f"date: {start or '2020-01-01'}",
    ]
    if authors:
        fm.append("authors:")
        fm.append(yaml_list(authors))
    fm.append("tags:")
    fm.extend(yaml_list(tags).splitlines())
    fm.append("external_link: ''")
    fm.append("image:")
    fm.append("  focal_point: Center")
    if code:
        fm.append(f"project_code: {yaml_quote(code)}")
    fm.append(f"project_status: {yaml_quote(as_str(row.get('status')))}")
    fm.append("---")
    fm.append("")
    return "\n".join(fm + body_lines) + "\n"


def render_patent(row: dict) -> str:
    title = as_str(row.get("title"))
    desc = as_str(row.get("description"))
    summary = desc[:100] + ("…" if len(desc) > 100 else "") if desc else title
    grant = as_date(row.get("grant_date"))
    app = as_date(row.get("application_date"))
    date = grant or app or "2020-01-01"
    number = as_str(row.get("patent_number"))
    inventors = as_str(row.get("inventors"))
    status = status_zh_patent(as_str(row.get("status")))
    authors = ["admin"] if "肖照林" in inventors else []
    tags = ["发明专利"]
    if as_str(row.get("status")).lower() == "granted":
        tags.append("已授权")

    body = [
        "**专利类型**：发明专利  ",
        f"**专利号**：{number or '—'}  ",
        f"**发明人**：{inventors or '—'}  ",
    ]
    if app:
        body.append(f"**申请日**：{app}  ")
    if grant:
        body.append(f"**授权日**：{grant}  ")
    body.append(f"**状态**：{status}  ")
    body.append("")
    if desc:
        body.append(desc)

    fm = [
        "---",
        f"title: {yaml_quote(title)}",
        f"summary: {yaml_quote(summary)}",
        f"date: {date}",
    ]
    if authors:
        fm.append("authors:")
        fm.append(yaml_list(authors))
    fm.append("tags:")
    fm.extend(yaml_list(tags).splitlines())
    if number:
        fm.append(f"patent_number: {yaml_quote(number)}")
    fm.append(f"patent_status: {yaml_quote(as_str(row.get('status')))}")
    fm.append("---")
    fm.append("")
    return "\n".join(fm + body) + "\n"


def ensure_section_index(section: Path, title: str, summary: str) -> None:
    section.mkdir(parents=True, exist_ok=True)
    idx = section / "_index.md"
    if not idx.exists():
        idx.write_text(
            f"---\ntitle: {title}\nview: card\nsummary: {summary}\n---\n",
            encoding="utf-8",
        )


def sync_section(
    *,
    label: str,
    section: Path,
    rows: list[dict],
    slug_fn,
    render_fn,
    dry_run: bool,
) -> dict[str, int]:
    """Upsert markdown pages; remove dirs not in Excel. Returns counts."""
    stats = {"added": 0, "updated": 0, "unchanged": 0, "removed": 0}
    wanted: set[str] = set()
    for row in rows:
        slug = slug_fn(row)
        wanted.add(slug)
        dest = section / slug / "index.md"
        content = render_fn(row)
        if dest.exists():
            old = dest.read_text(encoding="utf-8")
            if old == content:
                stats["unchanged"] += 1
                print(f"  = {label}/{slug}")
            else:
                stats["updated"] += 1
                print(f"  ~ {label}/{slug}")
                if not dry_run:
                    dest.write_text(content, encoding="utf-8")
        else:
            stats["added"] += 1
            print(f"  + {label}/{slug}")
            if not dry_run:
                dest.parent.mkdir(parents=True, exist_ok=True)
                dest.write_text(content, encoding="utf-8")

    for p in list(section.iterdir()) if section.exists() else []:
        if not p.is_dir():
            continue
        if p.name in wanted:
            continue
        stats["removed"] += 1
        print(f"  - {label}/{p.name}")
        if not dry_run:
            shutil.rmtree(p)
    return stats


def main() -> None:
    global PROJECT_XLSX, PATENT_XLSX, PROJECT_DIR, PATENT_DIR, CONTENT_DIR, XLSX_DIR
    ap = argparse.ArgumentParser(description="Import projects/patents from Excel")
    ap.add_argument(
        "--dry-run",
        action="store_true",
        help="Print add/update/remove summary without writing files",
    )
    ap.add_argument(
        "--xlsx-dir",
        type=Path,
        default=ROOT.parent,
        help="Directory containing 项目清单1.xlsx and 专利清单.xlsx",
    )
    ap.add_argument(
        "--project-xlsx",
        type=Path,
        default=None,
        help="Override project workbook path",
    )
    ap.add_argument(
        "--patent-xlsx",
        type=Path,
        default=None,
        help="Override patent workbook path",
    )
    ap.add_argument(
        "--content",
        type=Path,
        default=ROOT / "content",
        help="Hugo content directory (default: ./content)",
    )
    args = ap.parse_args()
    dry = args.dry_run
    XLSX_DIR = args.xlsx_dir.resolve()
    CONTENT_DIR = args.content.resolve()
    PROJECT_DIR = CONTENT_DIR / "project"
    PATENT_DIR = CONTENT_DIR / "patent"
    PROJECT_XLSX = (args.project_xlsx or (XLSX_DIR / "项目清单1.xlsx")).resolve()
    PATENT_XLSX = (args.patent_xlsx or (XLSX_DIR / "专利清单.xlsx")).resolve()

    if dry:
        print("DRY-RUN (no files written)")
    print(f"xlsx projects={PROJECT_XLSX}")
    print(f"xlsx patents={PATENT_XLSX}")
    print(f"content={CONTENT_DIR}")

    if not PROJECT_XLSX.exists():
        raise SystemExit(f"Missing {PROJECT_XLSX}")
    if not PATENT_XLSX.exists():
        raise SystemExit(f"Missing {PATENT_XLSX}")

    if not dry:
        ensure_section_index(PROJECT_DIR, "研究项目", "实验室承担的科研项目")
        ensure_section_index(PATENT_DIR, "专利", "实验室授权发明专利")

    print("projects:")
    ps = sync_section(
        label="project",
        section=PROJECT_DIR,
        rows=load_rows(PROJECT_XLSX),
        slug_fn=project_slug,
        render_fn=render_project,
        dry_run=dry,
    )
    print("patents:")
    pts = sync_section(
        label="patent",
        section=PATENT_DIR,
        rows=load_rows(PATENT_XLSX),
        slug_fn=patent_slug,
        render_fn=render_patent,
        dry_run=dry,
    )
    print(
        "summary projects "
        f"+{ps['added']} ~{ps['updated']} ={ps['unchanged']} -{ps['removed']} | "
        "patents "
        f"+{pts['added']} ~{pts['updated']} ={pts['unchanged']} -{pts['removed']}"
    )


if __name__ == "__main__":
    main()
